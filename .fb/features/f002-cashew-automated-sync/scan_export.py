#!/usr/bin/env python3
"""Scan a Cashew export (xlsx/csv) to ground f002 design decisions.

Reports: format, sheet(s), row count, column headers, a few sample rows, and
whether any column looks like a stable transaction ID (the thing f001's CSV
lacked). Run with the bench env python so openpyxl is available:

    /home/riz/work-bench/env/bin/python \
        apps/cashew_integration/.fb/features/f002-cashew-automated-sync/scan_export.py \
        /path/to/cashew-export.xlsx
"""
import sys, os, csv, json, re

ID_HINTS = ("id", "uuid", "guid", "key", "ref", "reference", "transaction_id", "pk")


def looks_like_id(header):
    h = header.strip().lower().replace(" ", "_")
    return any(re.fullmatch(rf"{hint}|.*_{hint}|{hint}_.*", h) for hint in ID_HINTS)


def report(headers, rows, sample_n=5):
    print(f"\nCOLUMNS ({len(headers)}):")
    for i, h in enumerate(headers):
        flag = "  <-- looks like an ID" if looks_like_id(str(h)) else ""
        print(f"  [{i}] {h!r}{flag}")

    id_cols = [h for h in headers if looks_like_id(str(h))]
    print(f"\nSTABLE-ID CANDIDATE COLUMNS: {id_cols or 'NONE (idempotency stays content-hash)'}")
    print(f"\nDATA ROWS: {len(rows)}")
    print(f"\nFIRST {min(sample_n, len(rows))} ROWS:")
    for r in rows[:sample_n]:
        print("  " + json.dumps({headers[i]: r[i] for i in range(min(len(headers), len(r)))}, default=str, ensure_ascii=False))


def scan_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"FORMAT: xlsx   SHEETS: {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    print(f"ACTIVE SHEET: {ws.title!r}   DIMENSIONS: {ws.dimensions}")
    grid = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    grid = [r for r in grid if any(v is not None and str(v).strip() != "" for v in r)]
    if not grid:
        print("EMPTY sheet."); return
    report([str(h) if h is not None else "" for h in grid[0]], grid[1:])


def scan_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    print(f"FORMAT: csv")
    rows = [r for r in rows if any(str(v).strip() for v in r)]
    if not rows:
        print("EMPTY file."); return
    report(rows[0], rows[1:])


def main():
    if len(sys.argv) < 2:
        sys.exit("usage: scan_export.py <path-to-export.(xlsx|csv)>")
    path = sys.argv[1]
    if not os.path.exists(path):
        sys.exit(f"no such file: {path}")
    print(f"FILE: {path}   SIZE: {os.path.getsize(path):,} bytes")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        scan_xlsx(path)
    elif ext == ".csv":
        scan_csv(path)
    else:
        sys.exit(f"unsupported extension {ext!r} — expected .xlsx or .csv")


if __name__ == "__main__":
    main()
